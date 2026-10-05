"""อ่าน PPT "ต่อสัญญา รถ MHE รอบประจำปี" → ข้อมูลตั้งต้นของแท็บ "รายละเอียดต่อสัญญา"
ใช้: python3 mhe-data/tools/proposal_extract.py <ไฟล์.pptx> > mhe-data/tools/proposal_base.json  (ต้องมี python-pptx)
ดึง: ราคาค่าเช่าต่อผู้ขาย (สไลด์ที่หัวตาราง "รถยกหมดสัญญา …") → รอบประมูล · Option ของ Supplier · After-sales Service · Standard Option (กลุ่มรูป)"""
import sys, re, json
from pptx import Presentation

def tx(s): return re.sub(r'\s+', ' ', (s or '').replace('\x0b', ' ')).strip()
def num(s):
    s = tx(s); n = s.replace(',', '')
    if re.fullmatch(r'\d+(\.\d+)?', n): f = float(n); return int(f) if f.is_integer() else f
    return s or None

def main(path):
    P = Presentation(path)
    vendors, rows, title, service, options, std = [], {}, '', {'suppliers': [], 'rows': []}, None, {}
    for s in P.slides:
        texts = [tx(sh.text_frame.text) for sh in s.shapes if sh.has_text_frame]
        for sh in s.shapes:
            if not sh.has_table: continue
            T = [[tx(c.text) for c in r.cells] for r in sh.table.rows]
            if T[0][0].startswith('รถยกหมดสัญญา'):                       # ตารางราคา
                title = T[0][0]; vc = list(range(3, len(T[1]), 2)); base = len(vendors)
                for j in vc: vendors.append({'name': T[1][j], 'contact': T[0][j]})
                for r in T[3:]:
                    if not r[0] or r[0] == 'รวม': continue
                    key = re.sub(r'\s*\(เดิม.*$', '', r[0]).lower(); row = rows.setdefault(key, {'item': r[0], 'energy': r[1], 'qty': num(r[2]), 'alloc': [], 'p': []})
                    while len(row['p']) < base: row['p'].append([None, None])
                    for j in vc: row['p'].append([num(r[j]), num(r[j + 1])])
            elif T[0][0].upper().startswith('SUPLLIER') or T[0][0].upper().startswith('SUPPLIER'):
                sup = T[0][1:]
                if any('ช่วงเวลาช่าง' in r[0] for r in T):                  # After-sales
                    off = len(service['suppliers']); service['suppliers'] += sup
                    for r in T[1:]:
                        lab = r[0].lstrip('- ').strip(); ex = next((x for x in service['rows'] if x[0] == lab), None)
                        if not ex: ex = [lab] + [''] * off; service['rows'].append(ex)
                        while len(ex) < off + 1: ex.append('')
                        ex += r[1:]
                else:                                                         # Option ของ Supplier (ช่องว่าง = รูปเครื่องหมายถูก)
                    options = {'suppliers': sup, 'rows': [[r[0]] + [c or '✓' for c in r[1:]] for r in T[1:]]}
        grp = [tx(g.text_frame.text) for sh in s.shapes if sh.shape_type == 6 for g in sh.shapes if g.has_text_frame and tx(g.text_frame.text)]
        head = next((g for g in grp if g.upper().startswith('ALL OPTION')), None)
        if head:
            items = [g for g in grp if g != head] + [t for t in texts if t and not re.match(r'(Duration|CCSS|Standards)', t)]
            std[head.replace('ALL OPTION', '').strip().title()] = items
    for r in rows.values():
        while len(r['p']) < len(vendors): r['p'].append([None, None])
    rnd = {'key': '2026-01', 'year': 2026, 'date': '2026-01-30', 'label': 'ม.ค. 2026', 'title': title,
           'sheets': [{'name': 'Lithium', 'vendors': vendors, 'rows': list(rows.values())}], 'confirm': []}
    return {'round': rnd, 'options': options, 'service': service, 'std': std,
            'note': '*** รถประเภทเดิมเป็น แบตน้ำ เปลี่ยนเป็นแบต Lithium ทั้งหมด'}

if __name__ == '__main__':
    print(json.dumps(main(sys.argv[1]), ensure_ascii=False, indent=1))
